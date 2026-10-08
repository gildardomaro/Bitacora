import openpyxl
import json
import datetime
import os

BOT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BOT_DIR, "scratch", "temp_nxt.xlsm")
BASE_CURVE_JSON = os.path.join(BOT_DIR, "real1000_base_curve.json")

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["REAL1000"]

curve = [{"time": "2026-05-14", "balance": 250.0}]

for r in range(6, ws.max_row + 1):
    f = ws.cell(r, 1).value
    bal = ws.cell(r, 15).value
    tin = ws.cell(r, 4).value
    if f is None and tin is None and bal is None:
        continue
    
    d = None
    if isinstance(f, (datetime.datetime, datetime.date)):
        d = f if isinstance(f, datetime.date) else f.date()
    elif isinstance(tin, (datetime.datetime, datetime.date)):
        d = tin if isinstance(tin, datetime.date) else tin.date()
    else:
        continue

    if bal is not None:
        curve.append({"time": str(d), "balance": round(float(bal), 2)})

with open(BASE_CURVE_JSON, "w", encoding="utf-8") as f:
    json.dump(curve, f, indent=2)

print(f"Base REAL1000 curve saved to {BASE_CURVE_JSON} with {len(curve)} points. Final bal: ${curve[-1]['balance']}")
