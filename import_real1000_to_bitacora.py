import os
import sys
import json
import shutil
import datetime
import openpyxl

BOT_DIR = os.path.dirname(os.path.abspath(__file__))
DASH_DIR = os.path.join(BOT_DIR, "dashboard")
EXCEL_SRC = os.path.join(BOT_DIR, "BITACORA-NXT-1000.xlsm")
TEMP_EXCEL = os.path.join(BOT_DIR, "scratch", "temp_import_nxt.xlsm")
HISTORICAL_JSON = os.path.join(BOT_DIR, "historical_real1000_trades.json")
STATIC_JS = os.path.join(BOT_DIR, "static_data.js")
DASH_STATIC_JS = os.path.join(DASH_DIR, "static_data.js")

def extract_trades_from_excel():
    print(f"Copiando Excel a {TEMP_EXCEL} para evitar bloqueos...")
    os.makedirs(os.path.dirname(TEMP_EXCEL), exist_ok=True)
    import subprocess
    cmd = f'powershell -Command "Copy-Item \'{EXCEL_SRC}\' \'{TEMP_EXCEL}\' -Force"'
    subprocess.run(cmd, shell=True, check=True)

    print("Cargando libro Excel...")
    wb = openpyxl.load_workbook(TEMP_EXCEL, data_only=True)
    ws = wb["REAL1000"]

    trades = []
    equity_curve = []
    
    # Deposito inicial fila 5: $250
    equity_curve.append({
        "time": "2026-05-14",
        "balance": 250.0
    })

    # Balance tracking
    month_initial_balances = {
        "2026-05": 250.0,
        "2026-06": 313.42,
        "2026-07": 452.07,
        "2026-08": 436.78,
        "2026-09": 499.59
    }

    for r in range(6, ws.max_row + 1):
        fecha_raw = ws.cell(r, 1).value
        inst_raw = ws.cell(r, 2).value
        tipo_raw = ws.cell(r, 3).value
        tin_raw = ws.cell(r, 4).value
        pin_raw = ws.cell(r, 5).value
        sl_raw = ws.cell(r, 6).value
        tp_raw = ws.cell(r, 7).value
        lot_raw = ws.cell(r, 8).value
        strat_raw = ws.cell(r, 9).value
        tout_raw = ws.cell(r, 10).value
        pout_raw = ws.cell(r, 11).value
        swap_raw = ws.cell(r, 12).value
        comm_raw = ws.cell(r, 13).value
        profit_raw = ws.cell(r, 14).value
        balance_raw = ws.cell(r, 15).value
        notas_raw = ws.cell(r, 16).value

        if fecha_raw is None and tin_raw is None:
            continue

        # Parse date
        if isinstance(fecha_raw, datetime.datetime):
            trade_date = fecha_raw.date()
        elif isinstance(fecha_raw, datetime.date):
            trade_date = fecha_raw
        elif isinstance(tin_raw, datetime.datetime):
            trade_date = tin_raw.date()
        else:
            continue

        # Parse time in
        if isinstance(tin_raw, datetime.datetime):
            dt_in = tin_raw
        elif isinstance(tin_raw, datetime.time):
            dt_in = datetime.datetime.combine(trade_date, tin_raw)
        else:
            dt_in = datetime.datetime.combine(trade_date, datetime.time(0, 0, 0))

        # Parse time out
        if isinstance(tout_raw, datetime.datetime):
            dt_out = tout_raw
        elif isinstance(tout_raw, datetime.time):
            if tout_raw < dt_in.time():
                dt_out = datetime.datetime.combine(trade_date + datetime.timedelta(days=1), tout_raw)
            else:
                dt_out = datetime.datetime.combine(trade_date, tout_raw)
        else:
            dt_out = dt_in

        # Month attribution: keep according to sheet structure (Row 204 belongs to August)
        month_str = trade_date.strftime("%Y-%m")
        if r <= 204 and month_str == "2026-09":
            month_str = "2026-08"

        profit = float(profit_raw) if profit_raw is not None else 0.0
        swap = float(swap_raw) if swap_raw is not None else 0.0
        comm = float(comm_raw) if comm_raw is not None else 0.0
        raw_profit = profit - (swap + comm)

        symbol = str(inst_raw or "GBPUSD").replace("/", "").strip().upper()
        order_type = str(tipo_raw or "BUY").strip().upper()
        lot = float(lot_raw) if lot_raw is not None else 0.02
        price_in = float(pin_raw) if pin_raw is not None else 0.0
        price_out = float(pout_raw) if pout_raw is not None else price_in

        # Extract ticket
        ticket_num = None
        if notas_raw and isinstance(notas_raw, str):
            import re
            m = re.search(r'\{?(\d{6,12})\}?', notas_raw)
            if m:
                ticket_num = int(m.group(1))
        if not ticket_num:
            ticket_num = 200000000 + r

        comment_str = ""
        if sl_raw and abs(price_out - float(sl_raw)) < 0.0005:
            comment_str = f"[sl {price_out:.5f}]"
        elif tp_raw and abs(price_out - float(tp_raw)) < 0.0005:
            comment_str = f"[tp {price_out:.5f}]"
        else:
            comment_str = f"REAL1000 #{ticket_num}"

        strategy_desc = str(strat_raw or "Manual (SMC + NexTrade)").strip()
        if "smart money" in strategy_desc.lower():
            strategy_desc = "Manual (SMC + NexTrade)"

        time_str = dt_out.strftime("%Y-%m-%d %H:%M:%S")
        date_str = dt_out.strftime("%Y-%m-%d")
        hour_str = dt_out.strftime("%H:%M:%S")

        t_dict = {
            "time": time_str,
            "date": date_str,
            "hour": hour_str,
            "month": month_str,
            "symbol": symbol,
            "type": order_type,
            "lot": lot,
            "volume": lot,
            "price": price_out,
            "price_open": price_in,
            "sl": float(sl_raw) if sl_raw is not None else 0.0,
            "tp": float(tp_raw) if tp_raw is not None else 0.0,
            "profit": round(profit, 2),
            "raw_profit": round(raw_profit, 2),
            "commission": round(comm, 2),
            "swap": round(swap, 2),
            "result": "WIN" if profit >= 0 else "LOSS",
            "ticket": ticket_num,
            "order": ticket_num,
            "account": "PEPPERSTONE",
            "terminal": "PEPPERSTONE",
            "strategy": strategy_desc,
            "comment": comment_str
        }
        trades.append(t_dict)

        if balance_raw is not None:
            equity_curve.append({
                "time": time_str,
                "balance": round(float(balance_raw), 2)
            })

    print(f"Extraídas exitosamente {len(trades)} operaciones de REAL1000.")
    with open(HISTORICAL_JSON, "w", encoding="utf-8") as f:
        json.dump(trades, f, indent=2, ensure_ascii=False)
    print(f"Historial guardado en {HISTORICAL_JSON}")

    return trades, month_initial_balances, equity_curve

def update_static_dashboard():
    hist_trades, month_deposits, real_equity_curve = extract_trades_from_excel()

    # Load current static_data.js
    if not os.path.exists(STATIC_JS):
        print(f"Error: {STATIC_JS} no encontrado.")
        return

    with open(STATIC_JS, "r", encoding="utf-8") as f:
        c = f.read()
    prefix = "window.STATIC_DASHBOARD_DATA ="
    json_str = c.split(prefix, 1)[1].strip()
    if json_str.endswith(";"):
        json_str = json_str[:-1].strip()
    data = json.loads(json_str)

    # 1. Separate current October trades from Pepperstone
    oct_trades = [t for t in data.get("real_trades", []) if t.get("month") == "2026-10"]
    if not oct_trades:
        # Fallback to existing trades if any
        oct_trades = [t for t in data.get("trades", []) if t.get("month") == "2026-10"]

    print(f"Operaciones actuales de Octubre en Pepperstone: {len(oct_trades)}")

    # 2. Combine all trades: historical (May-Sept) + October
    all_combined_trades = hist_trades + oct_trades
    all_combined_trades.sort(key=lambda x: x["time"])

    print(f"Total consolidado de operaciones para la cuenta activa: {len(all_combined_trades)}")

    # 3. Rebuild monthly analytics for PEPPERSTONE
    by_month = {}
    for t in all_combined_trades:
        m = t["month"]
        if m not in by_month:
            by_month[m] = {
                "trades": 0,
                "wins": 0,
                "losses": 0,
                "profit": 0.0,
                "commission_swap": 0.0,
                "list": []
            }
        p = t["profit"]
        c_sw = t["commission"] + t["swap"]
        by_month[m]["trades"] += 1
        by_month[m]["profit"] += p
        by_month[m]["commission_swap"] += c_sw
        if p >= 0:
            by_month[m]["wins"] += 1
        else:
            by_month[m]["losses"] += 1
        by_month[m]["list"].append(t)

    sorted_months = sorted(list(by_month.keys()))
    print("Meses disponibles en la cuenta:", sorted_months)

    # Referencias de Capital exactas de la tabla GRAFICO de la bitácora:
    # Mayo: $250, Junio: $500, Julio: $1,000, Agosto: $1,000, Septiembre: $2,000, Octubre: $3,000
    deposits_by_month = {
        "2026-05": 250.0,
        "2026-06": 500.0,
        "2026-07": 1000.0,
        "2026-08": 1000.0,
        "2026-09": 2000.0,
        "2026-10": 3000.0
    }

    monthly_series = []
    for m in sorted_months:
        inf = by_month[m]
        p = round(inf["profit"], 2)
        month_deposit = deposits_by_month.get(m, 1000.0)
        pct = round((p / month_deposit) * 100.0, 2)
        wr = round((inf["wins"] / inf["trades"]) * 100.0, 1) if inf["trades"] > 0 else 0.0

        monthly_series.append({
            "month": m,
            "profit_usd": p,
            "profit_pct": pct,
            "deposit_usd": month_deposit,
            "wins": inf["wins"],
            "losses": inf["losses"],
            "total_trades": inf["trades"],
            "win_rate": wr,
            "commission_swap_usd": round(inf["commission_swap"], 2)
        })

    total_acc_wins = sum(by_month[m]["wins"] for m in sorted_months)
    total_acc_losses = sum(by_month[m]["losses"] for m in sorted_months)
    total_acc_trades = sum(by_month[m]["trades"] for m in sorted_months)
    total_acc_profit = sum(by_month[m]["profit"] for m in sorted_months)
    global_wr = round((total_acc_wins / total_acc_trades) * 100.0, 1)

    print("\nResumen de la serie mensual consolidada:")
    for s in monthly_series:
        print(f"  {s['month']}: {s['profit_pct']:+6.2f}% | ${s['profit_usd']:+8.2f} USD | {s['wins']}G / {s['losses']}P (WR {s['win_rate']}%) [Base: ${s['deposit_usd']:.2f}]")

    print(f"\nTotales Globales: {total_acc_trades} trades | +${total_acc_profit:.2f} USD | WR {global_wr}%")

    # Update monthly_analytics
    if "monthly_analytics" not in data:
        data["monthly_analytics"] = {}

    data["monthly_analytics"]["PEPPERSTONE"] = {
        "terminal_key": "PEPPERSTONE",
        "name": "Pepperstone-01 (smart+IA y MARO)",
        "symbol": "XAUUSD / GBPUSD",
        "base_capital": 3000.0,
        "total_deposit_usd": 3000.0,
        "deposits_by_month": deposits_by_month,
        "total_wins": total_acc_wins,
        "total_losses": total_acc_losses,
        "total_trades": total_acc_trades,
        "win_rate": global_wr,
        "total_profit_usd": round(total_acc_profit, 2),
        "total_return_pct": round((total_acc_profit / 3000.0) * 100.0, 2),
        "monthly_series": monthly_series,
        "months_available": sorted_months,
        "current_month": "2026-10",
        "all_trades_by_month": {m: by_month[m]["list"] for m in sorted_months}
    }

    # Update real_trades and trades
    data["real_trades"] = all_combined_trades
    data["trades"] = all_combined_trades

    # Build equity curves
    equity_dict = {}

    # 1. Cuenta activa (Histórico REAL1000 + Octubre en vivo)
    combined_curve = list(real_equity_curve)
    curr_oct_bal = 3000.0
    combined_curve.append({"time": "2026-10-01", "balance": 3000.0})
    for t in oct_trades:
        curr_oct_bal += t["profit"]
        combined_curve.append({
            "time": t.get("time") or t["date"],
            "balance": round(curr_oct_bal, 2)
        })

    tot_wins = len([t for t in all_combined_trades if t["profit"] >= 0])
    gross_win = sum(t["profit"] for t in all_combined_trades if t["profit"] >= 0)
    gross_loss = sum(abs(t["profit"]) for t in all_combined_trades if t["profit"] < 0)
    pf_combined = round(gross_win / gross_loss, 2) if gross_loss > 0 else 1.78
    wr_combined = round(tot_wins / len(all_combined_trades) * 100.0, 1) if all_combined_trades else 65.8

    equity_dict["real1000"] = {
        "initial_balance": 250.0,
        "final_balance": round(curr_oct_bal, 2),
        "total_trades": len(all_combined_trades),
        "win_rate": wr_combined,
        "profit_factor": str(pf_combined),
        "curve": combined_curve
    }

    # 2. HECTOR
    hector_file = os.path.join(BOT_DIR, "hector_equity_data.json")
    if os.path.exists(hector_file):
        with open(hector_file, "r", encoding="utf-8") as hf:
            equity_dict["hector"] = json.load(hf)

    data["equity"] = equity_dict

    # 3. Radar Institucional y Calendario OANDA
    radar_cache = os.path.join(BOT_DIR, "radar_data_cache.json")
    if os.path.exists(radar_cache):
        try:
            with open(radar_cache, "r", encoding="utf-8") as rf:
                data["radar"] = json.load(rf)
        except Exception:
            pass
    else:
        try:
            from oanda_market_radar import generate_full_radar_payload
            data["radar"] = generate_full_radar_payload()
        except Exception:
            pass

    # Set timestamps
    now_utc = datetime.datetime.utcnow().isoformat() + "Z"
    now_local = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    data["generated_at"] = now_utc
    data["last_sync"] = now_local

    # Write output to static_data.js and dashboard/static_data.js
    output_content = "window.STATIC_DASHBOARD_DATA = " + json.dumps(data, indent=2, ensure_ascii=False) + ";\n"
    
    with open(STATIC_JS, "w", encoding="utf-8") as f:
        f.write(output_content)
    with open(DASH_STATIC_JS, "w", encoding="utf-8") as f:
        f.write(output_content)

    print(f"\nstatic_data.js actualizado con éxito en raíz ({STATIC_JS}) y dashboard ({DASH_STATIC_JS}).")

if __name__ == "__main__":
    update_static_dashboard()
